"""번역 결과를 마크다운과 Word(.docx)로 저장한다. 확인이 필요한 문단은 표시한다."""
from __future__ import annotations

from .pipeline import PipelineResult


def _md_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [[(c or "").replace("|", "\\|").replace("\n", " ") for c in r] + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(norm[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in norm[1:]]
    return "\n".join(lines)


def write_markdown(result: PipelineResult, path: str, title: str) -> None:
    parts = [f"# {title}", ""]
    for item in result.rendered():
        if item["kind"] == "table":
            parts.append(_md_table(item["tgt"]))
        else:
            parts.append(item["tgt"])
        if item["flagged"]:
            parts.append("<!-- 검수 필요: " + "; ".join(item["notes"]).replace("--", "-") + " -->")
        parts.append("")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(parts))


def write_docx(result: PipelineResult, path: str, title: str) -> None:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")

    doc.add_heading(title, level=1)
    for item in result.rendered():
        if item["kind"] == "table":
            rows = item["tgt"]
            if not rows:
                continue
            width = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    run = table.cell(r, c).paragraphs[0].add_run(cell or "")
                    if item["flagged"]:
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            doc.add_paragraph()
            continue
        p = doc.add_paragraph()
        run = p.add_run(item["tgt"])
        if item["flagged"]:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW   # 확인 필요 = 노란색
    doc.save(path)
